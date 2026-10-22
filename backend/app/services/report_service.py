import io
from datetime import datetime, timedelta, date
from typing import Optional, List, Dict, Any
from bson import ObjectId

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, landscape
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether, PageBreak
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfgen import canvas

from app.database.connection import db

# Collections
attendance_collection = db["attendance"]
leave_requests_collection = db["leave_requests"]
leave_types_collection = db["leave_types"]
leave_balances_collection = db["leave_balances"]
employees_collection = db["employees"]
departments_collection = db["departments"]
designations_collection = db["designations"]
meeting_rooms_collection = db["meeting_rooms"]
meeting_bookings_collection = db["meeting_bookings"]
workspace_desks_collection = db["workspace_desks"]
workspace_reservations_collection = db["workspace_reservations"]


# =========================================================================
# NUMBERED CANVAS FOR PROFESSIONAL PDF PAGINATION
# =========================================================================

class NumberedCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_decorations(num_pages)
            super().showPage()
        super().save()

    def draw_page_decorations(self, page_count):
        self.saveState()
        # Footer dividing line
        self.setStrokeColor(colors.HexColor("#CBD5E1"))
        self.setLineWidth(0.75)
        self.line(36, 32, self._pagesize[0] - 36, 32)

        # Footer text
        self.setFont("Helvetica-Bold", 8)
        self.setFillColor(colors.HexColor("#0F172A"))
        self.drawString(36, 18, "IntraSphere Enterprise")

        self.setFont("Helvetica", 8)
        self.setFillColor(colors.HexColor("#64748B"))
        self.drawString(135, 18, "• Smart Office Management System | Official System Report")

        page_str = f"Page {self._pageNumber} of {page_count}"
        self.drawRightString(self._pagesize[0] - 36, 18, page_str)
        self.restoreState()


# =========================================================================
# HELPER: DATE PARSING & RANGE NORMALIZATION
# =========================================================================

def parse_date_range(start_date_str: Optional[str] = None, end_date_str: Optional[str] = None):
    now = datetime.utcnow()
    if end_date_str and isinstance(end_date_str, str):
        try:
            end_dt = datetime.strptime(end_date_str, "%Y-%m-%d")
        except ValueError:
            end_dt = now
    else:
        end_dt = now

    if start_date_str and isinstance(start_date_str, str):
        try:
            start_dt = datetime.strptime(start_date_str, "%Y-%m-%d")
        except ValueError:
            start_dt = end_dt - timedelta(days=90)
    else:
        start_dt = end_dt - timedelta(days=90)

    start_dt = datetime.combine(start_dt.date(), datetime.min.time())
    end_dt = datetime.combine(end_dt.date(), datetime.max.time())

    # Count weekdays in range
    cur = start_dt.date()
    end_d = end_dt.date()
    working_days = 0
    total_days = max((end_d - cur).days + 1, 1)
    while cur <= end_d:
        if cur.weekday() < 5:  # Monday to Friday
            working_days += 1
        cur += timedelta(days=1)

    if working_days == 0:
        working_days = 1

    return start_dt, end_dt, working_days, total_days


# =========================================================================
# 1. ATTENDANCE REPORT DATA & EXPORTERS
# =========================================================================

async def get_attendance_report_data(
    start_date_str: Optional[str] = None,
    end_date_str: Optional[str] = None,
    department: Optional[str] = None,
    search: Optional[str] = None
) -> Dict[str, Any]:
    start_dt, end_dt, working_days, total_calendar_days = parse_date_range(start_date_str, end_date_str)

    # Department and Designation lookups
    depts = await departments_collection.find({}).to_list(length=200)
    dept_map = {str(d["_id"]): d["department_name"] for d in depts}
    desigs = await designations_collection.find({}).to_list(length=200)
    desig_map = {str(d["_id"]): d["designation_name"] for d in desigs}

    emp_query: Dict[str, Any] = {
        "$or": [
            {"is_active": True},
            {"employment_status": {"$in": ["Active", "active"]}},
            {"employment_status": {"$exists": False}}
        ]
    }
    if department and department != "All":
        dept_doc = await departments_collection.find_one({"department_name": department})
        if dept_doc:
            emp_query["department_id"] = dept_doc["_id"]

    employees = await employees_collection.find(emp_query).to_list(length=1000)

    if search:
        s_lower = search.lower()
        employees = [
            e for e in employees
            if s_lower in f"{e.get('first_name', '')} {e.get('last_name', '')}".lower()
            or s_lower in e.get("employee_id", "").lower()
            or s_lower in e.get("email", "").lower()
        ]

    employee_reports = []
    all_raw_logs = []

    total_present_all = 0
    total_late_all = 0
    total_ontime_all = 0
    total_early_all = 0
    total_hours_all = 0.0

    for emp in employees:
        emp_id = emp["_id"]
        emp_code = emp.get("employee_id", f"EMP-{str(emp_id)[-6:].upper()}")
        full_name = f"{emp.get('first_name', '')} {emp.get('last_name', '')}".strip() or "Unnamed Employee"
        email = emp.get("email", "")

        dept_val = emp.get("department")
        if isinstance(dept_val, str) and dept_val:
            dept_name = dept_val
        else:
            dept_name = dept_map.get(str(emp.get("department_id")), "General")

        desig_val = emp.get("designation")
        if isinstance(desig_val, str) and desig_val:
            desig_name = desig_val
        else:
            desig_name = desig_map.get(str(emp.get("designation_id")), "Staff")

        att_query: Dict[str, Any] = {
            "attendance_date": {"$gte": start_dt, "$lte": end_dt}
        }
        id_filters = [{"employee_id": emp_id}]
        if emp.get("user_id"):
            id_filters.append({"employee_id": emp["user_id"]})
        if emp_code:
            id_filters.append({"employee_code": emp_code})
        att_query["$or"] = id_filters

        records = await attendance_collection.find(att_query).sort("attendance_date", 1).to_list(length=500)

        total_days_present = len(records)
        on_time_days = 0
        late_days = 0
        early_checkout_days = 0
        emp_hours = 0.0
        total_check_in_minutes = 0

        for r in records:
            c_in = r.get("check_in")
            c_out = r.get("check_out")
            hours = float(r.get("working_hours") or 0.0)
            emp_hours += hours
            total_hours_all += hours
            total_present_all += 1

            is_late = False
            if c_in:
                total_check_in_minutes += c_in.hour * 60 + c_in.minute
                if c_in.hour > 9 or (c_in.hour == 9 and c_in.minute > 15):
                    late_days += 1
                    total_late_all += 1
                    is_late = True
                else:
                    on_time_days += 1
                    total_ontime_all += 1
            else:
                on_time_days += 1
                total_ontime_all += 1

            is_early = False
            if c_out and c_out.hour < 17:
                early_checkout_days += 1
                total_early_all += 1
                is_early = True

            att_d = r.get("attendance_date")
            d_str = att_d.strftime("%Y-%m-%d") if isinstance(att_d, (datetime, date)) else "N/A"
            in_str = c_in.strftime("%I:%M %p") if isinstance(c_in, datetime) else "N/A"
            out_str = c_out.strftime("%I:%M %p") if isinstance(c_out, datetime) else "In Progress"

            all_raw_logs.append({
                "employee_code": emp_code,
                "employee_name": full_name,
                "department": dept_name,
                "date": d_str,
                "check_in": in_str,
                "check_out": out_str,
                "hours": round(hours, 2),
                "status": "Late" if is_late else ("Early Leave" if is_early else "Present")
            })

        avg_daily_hours = round(emp_hours / total_days_present, 2) if total_days_present > 0 else 0.0
        punctuality_score = round((on_time_days / total_days_present) * 100, 1) if total_days_present > 0 else 100.0

        if total_days_present > 0 and total_check_in_minutes > 0:
            avg_min = total_check_in_minutes // total_days_present
            avg_h = avg_min // 60
            avg_m = avg_min % 60
            ampm = "AM" if avg_h < 12 else "PM"
            dis_h = avg_h if avg_h <= 12 else avg_h - 12
            if dis_h == 0:
                dis_h = 12
            avg_check_in_time = f"{dis_h:02d}:{avg_m:02d} {ampm}"
        else:
            avg_check_in_time = "N/A"

        if total_days_present == 0:
            grade = "No Data"
        elif punctuality_score >= 90:
            grade = "Excellent"
        elif punctuality_score >= 75:
            grade = "Good"
        elif punctuality_score >= 60:
            grade = "Average"
        else:
            grade = "Needs Improvement"

        employee_reports.append({
            "employee_code": emp_code,
            "employee_name": full_name,
            "email": email,
            "department": dept_name,
            "designation": desig_name,
            "total_days_present": total_days_present,
            "on_time_days": on_time_days,
            "late_days": late_days,
            "early_checkout_days": early_checkout_days,
            "avg_check_in_time": avg_check_in_time,
            "total_working_hours": round(emp_hours, 2),
            "avg_daily_hours": avg_daily_hours,
            "punctuality_score": punctuality_score,
            "grade": grade
        })

    total_possible_attendances = len(employees) * working_days
    overall_attendance_rate = round((total_present_all / max(total_possible_attendances, 1)) * 100, 1)
    overall_punctuality_rate = round((total_ontime_all / max(total_present_all, 1)) * 100, 1) if total_present_all > 0 else 100.0

    return {
        "metadata": {
            "report_title": "HR Administration Attendance & Punctuality Report",
            "start_date": start_dt.strftime("%Y-%m-%d"),
            "end_date": end_dt.strftime("%Y-%m-%d"),
            "working_days": working_days,
            "calendar_days": total_calendar_days,
            "department": department or "All Departments",
            "generated_at": datetime.utcnow().strftime("%Y-%m-%d %H:%M UTC"),
        },
        "summary": {
            "total_employees": len(employees),
            "total_present_instances": total_present_all,
            "total_ontime_instances": total_ontime_all,
            "total_late_instances": total_late_all,
            "total_early_instances": total_early_all,
            "total_hours_worked": round(total_hours_all, 2),
            "overall_attendance_rate": min(overall_attendance_rate, 100.0),
            "overall_punctuality_rate": min(overall_punctuality_rate, 100.0),
            "avg_daily_hours": round(total_hours_all / max(total_present_all, 1), 2)
        },
        "employees": employee_reports,
        "raw_logs": all_raw_logs
    }


def generate_attendance_excel(data: Dict[str, Any]) -> io.BytesIO:
    wb = Workbook()

    # Define color scheme (Modern Navy & IntraSphere Orange accent)
    NAVY = "0F172A"
    ORANGE = "EA580C"
    LIGHT_BG = "F8FAFC"
    BORDER_COLOR = "CBD5E1"

    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Segoe UI", size=15, bold=True, color=NAVY)
    subtitle_font = Font(name="Segoe UI", size=10, italic=True, color="475569")
    kpi_num_font = Font(name="Segoe UI", size=14, bold=True, color=ORANGE)
    kpi_label_font = Font(name="Segoe UI", size=9, bold=True, color="64748B")
    cell_font = Font(name="Segoe UI", size=10)
    bold_cell_font = Font(name="Segoe UI", size=10, bold=True)

    header_fill = PatternFill(start_color=NAVY, end_color=NAVY, fill_type="solid")
    kpi_fill = PatternFill(start_color=LIGHT_BG, end_color=LIGHT_BG, fill_type="solid")
    zebra_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color=BORDER_COLOR),
        right=Side(style='thin', color=BORDER_COLOR),
        top=Side(style='thin', color=BORDER_COLOR),
        bottom=Side(style='thin', color=BORDER_COLOR)
    )

    # -------------------------------------------------------------
    # SHEET 1: Attendance Summary & Punctuality
    # -------------------------------------------------------------
    ws1 = wb.active
    ws1.title = "Attendance Summary"
    ws1.views.sheetView[0].showGridLines = True

    meta = data["metadata"]
    summary = data["summary"]

    # Title Banner
    ws1.cell(row=1, column=1, value="INTRASPHERE ENTERPRISE").font = Font(name="Segoe UI", size=10, bold=True, color=ORANGE)
    ws1.cell(row=2, column=1, value="HR Attendance & Punctuality Evaluation Report").font = title_font
    ws1.cell(
        row=3, column=1,
        value=f"Period: {meta['start_date']} to {meta['end_date']}  |  Dept: {meta['department']}  |  Working Days: {meta['working_days']}  |  Generated: {meta['generated_at']}"
    ).font = subtitle_font

    # KPI Summary Cards Block (Row 5 - 6)
    kpis = [
        ("Total Staff", summary["total_employees"]),
        ("Total Present Logs", summary["total_present_instances"]),
        ("Total Late Arrivals", summary["total_late_instances"]),
        ("Attendance Rate", f"{summary['overall_attendance_rate']}%"),
        ("Punctuality Score", f"{summary['overall_punctuality_rate']}%"),
        ("Total Hours Worked", f"{summary['total_hours_worked']} hrs"),
        ("Avg Daily Hours", f"{summary['avg_daily_hours']} hrs"),
    ]

    for idx, (label, val) in enumerate(kpis, start=1):
        cell_l = ws1.cell(row=5, column=idx, value=label)
        cell_l.font = kpi_label_font
        cell_l.fill = kpi_fill
        cell_l.alignment = Alignment(horizontal="center", vertical="center")
        cell_l.border = thin_border

        cell_v = ws1.cell(row=6, column=idx, value=val)
        cell_v.font = kpi_num_font
        cell_v.fill = kpi_fill
        cell_v.alignment = Alignment(horizontal="center", vertical="center")
        cell_v.border = thin_border

    # Table Headers (Row 8)
    headers = [
        "Employee Code", "Employee Name", "Department", "Designation",
        "Days Present", "On-Time Days", "Late Days", "Early Departures",
        "Avg Check-In", "Total Hours", "Avg Hours/Day", "Punctuality %", "Grade"
    ]

    header_row = 8
    for col_idx, h in enumerate(headers, start=1):
        c = ws1.cell(row=header_row, column=col_idx, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center" if col_idx > 4 else "left", vertical="center")
        c.border = thin_border

    # Data Rows
    current_row = 9
    for emp in data["employees"]:
        row_vals = [
            emp["employee_code"],
            emp["employee_name"],
            emp["department"],
            emp["designation"],
            emp["total_days_present"],
            emp["on_time_days"],
            emp["late_days"],
            emp["early_checkout_days"],
            emp["avg_check_in_time"],
            emp["total_working_hours"],
            emp["avg_daily_hours"],
            f"{emp['punctuality_score']}%",
            emp["grade"]
        ]
        for col_idx, val in enumerate(row_vals, start=1):
            c = ws1.cell(row=current_row, column=col_idx, value=val)
            c.font = bold_cell_font if col_idx in (1, 12, 13) else cell_font
            c.alignment = Alignment(horizontal="center" if col_idx in (5, 6, 7, 8, 9, 10, 11, 12, 13) else "left", vertical="center")
            c.border = thin_border
            if (current_row % 2) == 0:
                c.fill = zebra_fill

            # Color highlight for Grade
            if col_idx == 13:
                if val == "Excellent":
                    c.font = Font(name="Segoe UI", size=10, bold=True, color="16A34A")
                elif val == "Needs Improvement":
                    c.font = Font(name="Segoe UI", size=10, bold=True, color="DC2626")
                elif val == "Good":
                    c.font = Font(name="Segoe UI", size=10, bold=True, color="2563EB")

        current_row += 1

    # Autosize columns
    for col in ws1.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws1.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # -------------------------------------------------------------
    # SHEET 2: Daily Logs Register
    # -------------------------------------------------------------
    ws2 = wb.create_sheet(title="Daily Check-In Logs")
    ws2.views.sheetView[0].showGridLines = True

    ws2.cell(row=1, column=1, value="INTRASPHERE ENTERPRISE").font = Font(name="Segoe UI", size=10, bold=True, color=ORANGE)
    ws2.cell(row=2, column=1, value="Detailed Daily Attendance Logs Register").font = title_font

    log_headers = ["Employee Code", "Employee Name", "Department", "Date", "Check-In", "Check-Out", "Duration (Hours)", "Status"]
    for col_idx, h in enumerate(log_headers, start=1):
        c = ws2.cell(row=4, column=col_idx, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center" if col_idx >= 4 else "left", vertical="center")
        c.border = thin_border

    log_row = 5
    for item in data["raw_logs"]:
        row_vals = [
            item["employee_code"],
            item["employee_name"],
            item["department"],
            item["date"],
            item["check_in"],
            item["check_out"],
            item["hours"],
            item["status"]
        ]
        for col_idx, val in enumerate(row_vals, start=1):
            c = ws2.cell(row=log_row, column=col_idx, value=val)
            c.font = cell_font
            c.alignment = Alignment(horizontal="center" if col_idx >= 4 else "left", vertical="center")
            c.border = thin_border
            if (log_row % 2) == 0:
                c.fill = zebra_fill
        log_row += 1

    for col in ws2.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws2.column_dimensions[col_letter].width = max(max_len + 4, 12)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def generate_attendance_pdf(data: Dict[str, Any]) -> io.BytesIO:
    buffer = io.BytesIO()
    # Landscape orientation for wide attendance report
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(letter),
        leftMargin=36,
        rightMargin=36,
        topMargin=36,
        bottomMargin=45
    )

    styles = getSampleStyleSheet()

    # Custom styles
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=18,
        leading=22,
        textColor=colors.HexColor("#0F172A")
    )
    subtitle_style = ParagraphStyle(
        'DocSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=13,
        textColor=colors.HexColor("#475569")
    )
    th_style = ParagraphStyle(
        'TableHeader',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8,
        leading=10,
        textColor=colors.white,
        alignment=1  # Center
    )
    td_style = ParagraphStyle(
        'TableCell',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=7.5,
        leading=10,
        textColor=colors.HexColor("#1E293B")
    )
    td_center_style = ParagraphStyle(
        'TableCellCenter',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=7.5,
        leading=10,
        textColor=colors.HexColor("#1E293B"),
        alignment=1
    )
    kpi_num_style = ParagraphStyle(
        'KpiNum',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=13,
        leading=16,
        textColor=colors.HexColor("#EA580C"),
        alignment=1
    )
    kpi_label_style = ParagraphStyle(
        'KpiLabel',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=7.5,
        leading=9,
        textColor=colors.HexColor("#475569"),
        alignment=1
    )

    meta = data["metadata"]
    summary = data["summary"]

    story = []

    # Brand & Header Banner
    story.append(Paragraph("<font color='#EA580C'><b>INTRASPHERE ENTERPRISE</b></font> • HR ADMINISTRATION", subtitle_style))
    story.append(Spacer(1, 3))
    story.append(Paragraph("Employee Attendance & Punctuality Official Report", title_style))
    story.append(Spacer(1, 4))
    story.append(Paragraph(
        f"<b>Filter Period:</b> {meta['start_date']} to {meta['end_date']} ({meta['working_days']} working days) &nbsp;|&nbsp; "
        f"<b>Department:</b> {meta['department']} &nbsp;|&nbsp; "
        f"<b>Generated On:</b> {meta['generated_at']}",
        subtitle_style
    ))
    story.append(Spacer(1, 14))

    # KPI Summary Cards Block
    kpi_data = [
        [
            Paragraph("TOTAL STAFF", kpi_label_style),
            Paragraph("TOTAL PRESENT", kpi_label_style),
            Paragraph("ON-TIME ARRIVALS", kpi_label_style),
            Paragraph("LATE ARRIVALS", kpi_label_style),
            Paragraph("ATTENDANCE RATE", kpi_label_style),
            Paragraph("PUNCTUALITY RATE", kpi_label_style),
            Paragraph("TOTAL HOURS", kpi_label_style),
        ],
        [
            Paragraph(str(summary["total_employees"]), kpi_num_style),
            Paragraph(str(summary["total_present_instances"]), kpi_num_style),
            Paragraph(str(summary["total_ontime_instances"]), kpi_num_style),
            Paragraph(str(summary["total_late_instances"]), kpi_num_style),
            Paragraph(f"{summary['overall_attendance_rate']}%", kpi_num_style),
            Paragraph(f"{summary['overall_punctuality_rate']}%", kpi_num_style),
            Paragraph(f"{summary['total_hours_worked']}h", kpi_num_style),
        ]
    ]
    kpi_table = Table(kpi_data, colWidths=[105] * 7)
    kpi_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor("#F8FAFC")),
        ('BOX', (0, 0), (-1, -1), 1, colors.HexColor("#E2E8F0")),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
    ]))
    story.append(kpi_table)
    story.append(Spacer(1, 16))

    # Main Attendance Table
    table_headers = [
        Paragraph("Code", th_style),
        Paragraph("Employee Name", th_style),
        Paragraph("Department", th_style),
        Paragraph("Designation", th_style),
        Paragraph("Present", th_style),
        Paragraph("On-Time", th_style),
        Paragraph("Late", th_style),
        Paragraph("Early", th_style),
        Paragraph("Avg Check-In", th_style),
        Paragraph("Total Hrs", th_style),
        Paragraph("Avg Hrs/Day", th_style),
        Paragraph("Punctuality", th_style),
        Paragraph("Grade", th_style),
    ]

    col_widths = [55, 110, 75, 75, 42, 45, 38, 38, 65, 52, 55, 55, 65]
    table_data = [table_headers]

    for emp in data["employees"]:
        grade_color = "#16A34A" if emp["grade"] == "Excellent" else ("#DC2626" if emp["grade"] == "Needs Improvement" else "#2563EB")
        table_data.append([
            Paragraph(emp["employee_code"], td_style),
            Paragraph(emp["employee_name"], td_style),
            Paragraph(emp["department"], td_style),
            Paragraph(emp["designation"], td_style),
            Paragraph(str(emp["total_days_present"]), td_center_style),
            Paragraph(str(emp["on_time_days"]), td_center_style),
            Paragraph(str(emp["late_days"]), td_center_style),
            Paragraph(str(emp["early_checkout_days"]), td_center_style),
            Paragraph(emp["avg_check_in_time"], td_center_style),
            Paragraph(f"{emp['total_working_hours']}h", td_center_style),
            Paragraph(f"{emp['avg_daily_hours']}h", td_center_style),
            Paragraph(f"{emp['punctuality_score']}%", td_center_style),
            Paragraph(f"<font color='{grade_color}'><b>{emp['grade']}</b></font>", td_center_style),
        ])

    main_table = Table(table_data, colWidths=col_widths, repeatRows=1)
    main_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#0F172A")),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
    ]))
    story.append(main_table)

    doc.build(story, canvasmaker=NumberedCanvas)
    buffer.seek(0)
    return buffer


# =========================================================================
# 2. LEAVE REPORT DATA & EXPORTERS
# =========================================================================

async def get_leave_report_data(
    start_date_str: Optional[str] = None,
    end_date_str: Optional[str] = None,
    department: Optional[str] = None,
    status_filter: Optional[str] = None,
    leave_type_id: Optional[str] = None,
    search: Optional[str] = None
) -> Dict[str, Any]:
    start_dt, end_dt, working_days, total_calendar_days = parse_date_range(start_date_str, end_date_str)

    # Leave Types lookup
    ltypes = await leave_types_collection.find({}).to_list(length=100)
    ltype_map = {str(lt["_id"]): lt["name"] for lt in ltypes}

    # Employees lookup (index by both _id and user_id)
    emps = await employees_collection.find({}).to_list(length=1000)
    emp_map = {}
    for e in emps:
        emp_map[str(e["_id"])] = e
        if e.get("user_id"):
            emp_map[str(e["user_id"])] = e

    query: Dict[str, Any] = {
        "start_date": {"$gte": start_dt, "$lte": end_dt}
    }
    if status_filter and status_filter != "All":
        query["status"] = status_filter
    if leave_type_id and leave_type_id != "All":
        try:
            query["leave_type_id"] = ObjectId(leave_type_id)
        except Exception:
            pass

    records = await leave_requests_collection.find(query).sort("created_at", -1).to_list(length=2000)

    # Department and search filter
    filtered_records = []
    total_approved = 0
    total_pending = 0
    total_rejected = 0
    total_cancelled = 0
    total_days_approved = 0.0
    type_counts: Dict[str, int] = {}

    for req in records:
        emp_id = str(req.get("employee_id"))
        emp = emp_map.get(emp_id)
        emp_dept = emp.get("department", "General") if emp else "General"

        if department and department != "All" and emp_dept != department:
            continue

        if search:
            s_lower = search.lower()
            if emp:
                e_full = f"{emp.get('first_name', '')} {emp.get('last_name', '')}".lower()
                e_mail = emp.get("email", "").lower()
                e_code = emp.get("employee_id", "").lower()
                if s_lower not in e_full and s_lower not in e_mail and s_lower not in e_code:
                    continue
            else:
                continue

        st = req.get("status", "Pending")
        days = float(req.get("total_days", 1.0))

        if st == "Approved":
            total_approved += 1
            total_days_approved += days
        elif st == "Pending":
            total_pending += 1
        elif st == "Rejected":
            total_rejected += 1
        elif st in ("Withdrawn", "Cancelled"):
            total_cancelled += 1

        lt_id = str(req.get("leave_type_id"))
        lt_name = ltype_map.get(lt_id, "Standard Leave")
        type_counts[lt_name] = type_counts.get(lt_name, 0) + 1

        emp_code = emp.get("employee_id", "EMP-000") if emp else "EMP-000"
        emp_name = f"{emp.get('first_name', '')} {emp.get('last_name', '')}".strip() if emp else "Unknown Employee"

        s_date = req["start_date"].strftime("%Y-%m-%d") if isinstance(req.get("start_date"), (datetime, date)) else "N/A"
        e_date = req["end_date"].strftime("%Y-%m-%d") if isinstance(req.get("end_date"), (datetime, date)) else "N/A"

        filtered_records.append({
            "request_id": str(req["_id"]),
            "employee_code": emp_code,
            "employee_name": emp_name,
            "department": emp_dept,
            "leave_type": lt_name,
            "start_date": s_date,
            "end_date": e_date,
            "total_days": days,
            "is_half_day": "Yes" if req.get("is_half_day") else "No",
            "half_day_session": req.get("half_day_session") or "-",
            "status": st,
            "reason": req.get("reason", "N/A"),
            "approval_comment": req.get("approval_comment") or req.get("rejection_reason") or "-",
            "created_at": req.get("created_at").strftime("%Y-%m-%d %H:%M") if isinstance(req.get("created_at"), datetime) else "-"
        })

    # Leave balances for employee register
    balances_data = []
    cursor_bal = leave_balances_collection.find({"year": 2026}).sort("employee_id", 1)
    async for b in cursor_bal:
        emp = emp_map.get(str(b["employee_id"]))
        if not emp:
            continue
        emp_dept = emp.get("department", "General")
        if department and department != "All" and emp_dept != department:
            continue

        emp_code = emp.get("employee_id", "EMP-000")
        emp_name = f"{emp.get('first_name', '')} {emp.get('last_name', '')}".strip()
        lt_name = ltype_map.get(str(b["leave_type_id"]), "Leave")

        balances_data.append({
            "employee_code": emp_code,
            "employee_name": emp_name,
            "department": emp_dept,
            "leave_type": lt_name,
            "allocated": float(b.get("allocated", 0)),
            "used": float(b.get("used", 0)),
            "pending": float(b.get("pending", 0)),
            "remaining": float(b.get("remaining", 0))
        })

    return {
        "metadata": {
            "report_title": "HR Administration Leave Management Report",
            "start_date": start_dt.strftime("%Y-%m-%d"),
            "end_date": end_dt.strftime("%Y-%m-%d"),
            "department": department or "All Departments",
            "status_filter": status_filter or "All",
            "generated_at": datetime.utcnow().strftime("%Y-%m-%d %H:%M UTC")
        },
        "summary": {
            "total_requests": len(filtered_records),
            "approved": total_approved,
            "pending": total_pending,
            "rejected": total_rejected,
            "cancelled": total_cancelled,
            "total_days_approved": round(total_days_approved, 1),
            "type_counts": type_counts
        },
        "requests": filtered_records,
        "balances": balances_data
    }


def generate_leave_excel(data: Dict[str, Any]) -> io.BytesIO:
    wb = Workbook()

    NAVY = "0F172A"
    ORANGE = "EA580C"
    LIGHT_BG = "F8FAFC"
    BORDER_COLOR = "CBD5E1"

    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Segoe UI", size=15, bold=True, color=NAVY)
    subtitle_font = Font(name="Segoe UI", size=10, italic=True, color="475569")
    kpi_num_font = Font(name="Segoe UI", size=14, bold=True, color=ORANGE)
    kpi_label_font = Font(name="Segoe UI", size=9, bold=True, color="64748B")
    cell_font = Font(name="Segoe UI", size=10)
    bold_cell_font = Font(name="Segoe UI", size=10, bold=True)

    header_fill = PatternFill(start_color=NAVY, end_color=NAVY, fill_type="solid")
    kpi_fill = PatternFill(start_color=LIGHT_BG, end_color=LIGHT_BG, fill_type="solid")
    zebra_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color=BORDER_COLOR),
        right=Side(style='thin', color=BORDER_COLOR),
        top=Side(style='thin', color=BORDER_COLOR),
        bottom=Side(style='thin', color=BORDER_COLOR)
    )

    # SHEET 1: Leave Applications Register
    ws1 = wb.active
    ws1.title = "Leave Applications"
    ws1.views.sheetView[0].showGridLines = True

    meta = data["metadata"]
    summary = data["summary"]

    ws1.cell(row=1, column=1, value="INTRASPHERE ENTERPRISE").font = Font(name="Segoe UI", size=10, bold=True, color=ORANGE)
    ws1.cell(row=2, column=1, value="HR Employee Leave Management & Audit Report").font = title_font
    ws1.cell(
        row=3, column=1,
        value=f"Period: {meta['start_date']} to {meta['end_date']}  |  Dept: {meta['department']}  |  Status: {meta['status_filter']}  |  Generated: {meta['generated_at']}"
    ).font = subtitle_font

    # KPI Block
    kpis = [
        ("Total Requests", summary["total_requests"]),
        ("Approved Requests", summary["approved"]),
        ("Pending Approvals", summary["pending"]),
        ("Rejected Requests", summary["rejected"]),
        ("Withdrawn / Cancelled", summary["cancelled"]),
        ("Total Days Approved", f"{summary['total_days_approved']} Days"),
    ]
    for idx, (label, val) in enumerate(kpis, start=1):
        cell_l = ws1.cell(row=5, column=idx, value=label)
        cell_l.font = kpi_label_font
        cell_l.fill = kpi_fill
        cell_l.alignment = Alignment(horizontal="center", vertical="center")
        cell_l.border = thin_border

        cell_v = ws1.cell(row=6, column=idx, value=val)
        cell_v.font = kpi_num_font
        cell_v.fill = kpi_fill
        cell_v.alignment = Alignment(horizontal="center", vertical="center")
        cell_v.border = thin_border

    # Headers
    headers = [
        "Employee Code", "Employee Name", "Department", "Leave Type",
        "Start Date", "End Date", "Total Days", "Half-Day?", "Status", "Reason", "Notes / Feedback", "Applied On"
    ]
    header_row = 8
    for col_idx, h in enumerate(headers, start=1):
        c = ws1.cell(row=header_row, column=col_idx, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center" if col_idx in (5, 6, 7, 8, 9, 12) else "left", vertical="center")
        c.border = thin_border

    # Rows
    current_row = 9
    for req in data["requests"]:
        row_vals = [
            req["employee_code"],
            req["employee_name"],
            req["department"],
            req["leave_type"],
            req["start_date"],
            req["end_date"],
            req["total_days"],
            f"{req['is_half_day']} ({req['half_day_session']})" if req["is_half_day"] == "Yes" else "No",
            req["status"],
            req["reason"],
            req["approval_comment"],
            req["created_at"]
        ]
        for col_idx, val in enumerate(row_vals, start=1):
            c = ws1.cell(row=current_row, column=col_idx, value=val)
            c.font = bold_cell_font if col_idx in (1, 9) else cell_font
            c.alignment = Alignment(horizontal="center" if col_idx in (5, 6, 7, 8, 9, 12) else "left", vertical="center")
            c.border = thin_border
            if (current_row % 2) == 0:
                c.fill = zebra_fill

            if col_idx == 9:
                if val == "Approved":
                    c.font = Font(name="Segoe UI", size=10, bold=True, color="16A34A")
                elif val == "Rejected":
                    c.font = Font(name="Segoe UI", size=10, bold=True, color="DC2626")
                elif val == "Pending":
                    c.font = Font(name="Segoe UI", size=10, bold=True, color="D97706")

        current_row += 1

    for col in ws1.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws1.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # SHEET 2: Employee Leave Balances
    ws2 = wb.create_sheet(title="Leave Balances")
    ws2.views.sheetView[0].showGridLines = True

    ws2.cell(row=1, column=1, value="INTRASPHERE ENTERPRISE").font = Font(name="Segoe UI", size=10, bold=True, color=ORANGE)
    ws2.cell(row=2, column=1, value="Annual Employee Leave Quota & Remaining Balances").font = title_font

    bal_headers = ["Employee Code", "Employee Name", "Department", "Leave Type", "Allocated (Days)", "Used (Days)", "Pending (Days)", "Remaining Balance"]
    for col_idx, h in enumerate(bal_headers, start=1):
        c = ws2.cell(row=4, column=col_idx, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center" if col_idx >= 5 else "left", vertical="center")
        c.border = thin_border

    bal_row = 5
    for b in data["balances"]:
        row_vals = [
            b["employee_code"],
            b["employee_name"],
            b["department"],
            b["leave_type"],
            b["allocated"],
            b["used"],
            b["pending"],
            b["remaining"]
        ]
        for col_idx, val in enumerate(row_vals, start=1):
            c = ws2.cell(row=bal_row, column=col_idx, value=val)
            c.font = bold_cell_font if col_idx == 8 else cell_font
            c.alignment = Alignment(horizontal="center" if col_idx >= 5 else "left", vertical="center")
            c.border = thin_border
            if (bal_row % 2) == 0:
                c.fill = zebra_fill
        bal_row += 1

    for col in ws2.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws2.column_dimensions[col_letter].width = max(max_len + 4, 12)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def generate_leave_pdf(data: Dict[str, Any]) -> io.BytesIO:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(letter),
        leftMargin=36,
        rightMargin=36,
        topMargin=36,
        bottomMargin=45
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=18,
        leading=22,
        textColor=colors.HexColor("#0F172A")
    )
    subtitle_style = ParagraphStyle(
        'DocSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=13,
        textColor=colors.HexColor("#475569")
    )
    th_style = ParagraphStyle(
        'TableHeader',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8,
        leading=10,
        textColor=colors.white,
        alignment=1
    )
    td_style = ParagraphStyle(
        'TableCell',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=7.5,
        leading=10,
        textColor=colors.HexColor("#1E293B")
    )
    td_center_style = ParagraphStyle(
        'TableCellCenter',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=7.5,
        leading=10,
        textColor=colors.HexColor("#1E293B"),
        alignment=1
    )
    kpi_num_style = ParagraphStyle(
        'KpiNum',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=13,
        leading=16,
        textColor=colors.HexColor("#EA580C"),
        alignment=1
    )
    kpi_label_style = ParagraphStyle(
        'KpiLabel',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=7.5,
        leading=9,
        textColor=colors.HexColor("#475569"),
        alignment=1
    )

    meta = data["metadata"]
    summary = data["summary"]

    story = []
    story.append(Paragraph("<font color='#EA580C'><b>INTRASPHERE ENTERPRISE</b></font> • HR ADMINISTRATION", subtitle_style))
    story.append(Spacer(1, 3))
    story.append(Paragraph("Leave Management & Employee Absence Report", title_style))
    story.append(Spacer(1, 4))
    story.append(Paragraph(
        f"<b>Filter Period:</b> {meta['start_date']} to {meta['end_date']} &nbsp;|&nbsp; "
        f"<b>Department:</b> {meta['department']} &nbsp;|&nbsp; "
        f"<b>Status Filter:</b> {meta['status_filter']} &nbsp;|&nbsp; "
        f"<b>Generated On:</b> {meta['generated_at']}",
        subtitle_style
    ))
    story.append(Spacer(1, 14))

    # KPI Table
    kpi_data = [
        [
            Paragraph("TOTAL APPLICATIONS", kpi_label_style),
            Paragraph("APPROVED LEAVES", kpi_label_style),
            Paragraph("PENDING APPROVALS", kpi_label_style),
            Paragraph("REJECTED LEAVES", kpi_label_style),
            Paragraph("WITHDRAWN / CANCELLED", kpi_label_style),
            Paragraph("TOTAL DAYS APPROVED", kpi_label_style),
        ],
        [
            Paragraph(str(summary["total_requests"]), kpi_num_style),
            Paragraph(str(summary["approved"]), kpi_num_style),
            Paragraph(str(summary["pending"]), kpi_num_style),
            Paragraph(str(summary["rejected"]), kpi_num_style),
            Paragraph(str(summary["cancelled"]), kpi_num_style),
            Paragraph(f"{summary['total_days_approved']} Days", kpi_num_style),
        ]
    ]
    kpi_table = Table(kpi_data, colWidths=[122] * 6)
    kpi_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor("#F8FAFC")),
        ('BOX', (0, 0), (-1, -1), 1, colors.HexColor("#E2E8F0")),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
    ]))
    story.append(kpi_table)
    story.append(Spacer(1, 16))

    # Leave Table
    table_headers = [
        Paragraph("Code", th_style),
        Paragraph("Employee Name", th_style),
        Paragraph("Department", th_style),
        Paragraph("Leave Type", th_style),
        Paragraph("Start Date", th_style),
        Paragraph("End Date", th_style),
        Paragraph("Days", th_style),
        Paragraph("Status", th_style),
        Paragraph("Reason", th_style),
        Paragraph("Manager Comments", th_style),
    ]

    col_widths = [60, 115, 80, 85, 65, 65, 45, 65, 120, 95]
    table_data = [table_headers]

    for req in data["requests"]:
        st = req["status"]
        st_color = "#16A34A" if st == "Approved" else ("#DC2626" if st == "Rejected" else "#D97706")
        table_data.append([
            Paragraph(req["employee_code"], td_style),
            Paragraph(req["employee_name"], td_style),
            Paragraph(req["department"], td_style),
            Paragraph(req["leave_type"], td_style),
            Paragraph(req["start_date"], td_center_style),
            Paragraph(req["end_date"], td_center_style),
            Paragraph(str(req["total_days"]), td_center_style),
            Paragraph(f"<font color='{st_color}'><b>{st}</b></font>", td_center_style),
            Paragraph(req["reason"][:45] + ("..." if len(req["reason"]) > 45 else ""), td_style),
            Paragraph(req["approval_comment"][:35] + ("..." if len(req["approval_comment"]) > 35 else ""), td_style),
        ])

    main_table = Table(table_data, colWidths=col_widths, repeatRows=1)
    main_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#0F172A")),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
    ]))
    story.append(main_table)

    doc.build(story, canvasmaker=NumberedCanvas)
    buffer.seek(0)
    return buffer


# =========================================================================
# 3. FACILITY MANAGER: RESOURCE UTILIZATION REPORT DATA & EXPORTERS
# =========================================================================

async def get_resource_utilization_report_data(
    start_date_str: Optional[str] = None,
    end_date_str: Optional[str] = None,
    floor: Optional[int] = None,
    zone: Optional[str] = None,
    resource_type: Optional[str] = None  # 'rooms', 'workspaces', 'all'
) -> Dict[str, Any]:
    start_dt, end_dt, working_days, total_calendar_days = parse_date_range(start_date_str, end_date_str)

    # Standard office hours = 9h per business day
    DAILY_OPERATING_HOURS = 9.0
    total_operating_hours_per_unit = working_days * DAILY_OPERATING_HOURS

    # -------------------------------------------------------------
    # A. Meeting Rooms Occupancy Aggregation
    # -------------------------------------------------------------
    room_query: Dict[str, Any] = {"is_active": True}
    if floor is not None:
        room_query["floor"] = floor

    all_rooms = await meeting_rooms_collection.find(room_query).to_list(length=200)

    # Fetch confirmed/completed bookings in range
    booking_query: Dict[str, Any] = {
        "status": {"$in": ["Confirmed", "Completed"]},
        "start_time": {"$gte": start_dt, "$lte": end_dt}
    }
    all_bookings = await meeting_bookings_collection.find(booking_query).to_list(length=2000)

    # Hourly distribution for peak identification
    hourly_room_distribution: Dict[int, int] = {h: 0 for h in range(8, 20)}

    room_reports = []
    total_room_booked_hours = 0.0
    total_room_capacity = 0

    for rm in all_rooms:
        rm_id = rm["_id"]
        rm_bookings = [b for b in all_bookings if b.get("room_id") == rm_id]

        booked_hours = 0.0
        for b in rm_bookings:
            st = b.get("start_time")
            et = b.get("end_time")
            if isinstance(st, datetime) and isinstance(et, datetime) and et > st:
                diff_hours = (et - st).total_seconds() / 3600.0
                booked_hours += diff_hours
                # Record hour histogram
                for h in range(st.hour, min(et.hour + 1, 20)):
                    if h in hourly_room_distribution:
                        hourly_room_distribution[h] += 1

        occupancy_rate = min(100.0, round((booked_hours / max(total_operating_hours_per_unit, 1.0)) * 100, 1))
        total_room_booked_hours += booked_hours
        cap = int(rm.get("capacity", 4))
        total_room_capacity += cap

        room_reports.append({
            "room_id": str(rm_id),
            "room_name": rm.get("room_name", "Meeting Room"),
            "room_code": rm.get("room_code", "CR-000"),
            "floor": rm.get("floor", 1),
            "location": rm.get("location", "Main Office"),
            "capacity": cap,
            "status": rm.get("status", "Available"),
            "total_bookings": len(rm_bookings),
            "total_hours_booked": round(booked_hours, 1),
            "operating_hours": total_operating_hours_per_unit,
            "occupancy_rate_percent": occupancy_rate
        })

    # Sort rooms by occupancy rate descending
    room_reports.sort(key=lambda x: x["occupancy_rate_percent"], reverse=True)

    total_possible_room_hours = len(all_rooms) * total_operating_hours_per_unit
    avg_room_occupancy = round((total_room_booked_hours / max(total_possible_room_hours, 1.0)) * 100, 1)

    # -------------------------------------------------------------
    # B. Workspaces & Desks Utilization Aggregation
    # -------------------------------------------------------------
    desk_query: Dict[str, Any] = {"is_active": True}
    if floor is not None:
        desk_query["floor"] = floor
    if zone and zone != "All":
        desk_query["zone"] = zone

    all_desks = await workspace_desks_collection.find(desk_query).to_list(length=500)

    # Workspace reservations
    desk_res_query: Dict[str, Any] = {
        "status": {"$in": ["Confirmed", "Completed"]},
        "start_time": {"$gte": start_dt, "$lte": end_dt}
    }
    all_desk_res = await workspace_reservations_collection.find(desk_res_query).to_list(length=3000)

    hourly_desk_distribution: Dict[int, int] = {h: 0 for h in range(8, 20)}

    desk_reports = []
    zone_stats: Dict[str, Dict[str, Any]] = {}
    total_desk_reserved_hours = 0.0

    for dk in all_desks:
        dk_id = dk["_id"]
        dk_res = [r for r in all_desk_res if r.get("desk_id") == dk_id]

        reserved_hours = 0.0
        for r in dk_res:
            st = r.get("start_time")
            et = r.get("end_time")
            if isinstance(st, datetime) and isinstance(et, datetime) and et > st:
                diff_hours = (et - st).total_seconds() / 3600.0
                reserved_hours += diff_hours
                for h in range(st.hour, min(et.hour + 1, 20)):
                    if h in hourly_desk_distribution:
                        hourly_desk_distribution[h] += 1

        utilization_rate = min(100.0, round((reserved_hours / max(total_operating_hours_per_unit, 1.0)) * 100, 1))
        total_desk_reserved_hours += reserved_hours

        z_name = dk.get("zone", "General Open Workspace")
        if z_name not in zone_stats:
            zone_stats[z_name] = {
                "zone_name": z_name,
                "floor": dk.get("floor", 1),
                "total_desks": 0,
                "total_reservations": 0,
                "total_hours": 0.0
            }
        zone_stats[z_name]["total_desks"] += 1
        zone_stats[z_name]["total_reservations"] += len(dk_res)
        zone_stats[z_name]["total_hours"] += reserved_hours

        desk_reports.append({
            "desk_id": str(dk_id),
            "desk_code": dk.get("desk_code", "D-000"),
            "desk_name": dk.get("desk_name", "Hot Desk"),
            "floor": dk.get("floor", 1),
            "zone": z_name,
            "workspace_type": dk.get("workspace_type", "Standard Desk"),
            "status": dk.get("status", "Available"),
            "total_reservations": len(dk_res),
            "total_hours_reserved": round(reserved_hours, 1),
            "operating_hours": total_operating_hours_per_unit,
            "utilization_rate_percent": utilization_rate
        })

    desk_reports.sort(key=lambda x: x["utilization_rate_percent"], reverse=True)

    # Compute zone averages
    zone_reports = []
    for z_name, z in zone_stats.items():
        z_possible_hours = z["total_desks"] * total_operating_hours_per_unit
        z_util = round((z["total_hours"] / max(z_possible_hours, 1.0)) * 100, 1)
        zone_reports.append({
            "zone_name": z_name,
            "floor": z["floor"],
            "total_desks": z["total_desks"],
            "total_reservations": z["total_reservations"],
            "total_hours_reserved": round(z["total_hours"], 1),
            "utilization_rate_percent": min(z_util, 100.0)
        })
    zone_reports.sort(key=lambda x: x["utilization_rate_percent"], reverse=True)

    total_possible_desk_hours = len(all_desks) * total_operating_hours_per_unit
    avg_desk_utilization = round((total_desk_reserved_hours / max(total_possible_desk_hours, 1.0)) * 100, 1)

    # Find peak hours
    peak_room_hour = max(hourly_room_distribution, key=hourly_room_distribution.get) if hourly_room_distribution else 10
    peak_desk_hour = max(hourly_desk_distribution, key=hourly_desk_distribution.get) if hourly_desk_distribution else 11

    def fmt_hour_range(h: int):
        ampm1 = "AM" if h < 12 else "PM"
        ampm2 = "AM" if (h + 1) < 12 else "PM"
        h1 = h if h <= 12 else h - 12
        h2 = (h + 1) if (h + 1) <= 12 else (h + 1) - 12
        return f"{h1:02d}:00 {ampm1} - {h2:02d}:00 {ampm2}"

    # Overall Facility Space Efficiency score (weighted 50% room, 50% workspace)
    overall_efficiency = round((min(avg_room_occupancy, 100.0) * 0.5) + (min(avg_desk_utilization, 100.0) * 0.5), 1)

    return {
        "metadata": {
            "report_title": "Facility Resource Utilization & Occupancy Report",
            "start_date": start_dt.strftime("%Y-%m-%d"),
            "end_date": end_dt.strftime("%Y-%m-%d"),
            "working_days": working_days,
            "operating_hours_per_unit": total_operating_hours_per_unit,
            "floor": f"Floor {floor}" if floor is not None else "All Floors",
            "zone": zone or "All Zones",
            "generated_at": datetime.utcnow().strftime("%Y-%m-%d %H:%M UTC")
        },
        "summary": {
            "overall_space_efficiency_score": overall_efficiency,
            "meeting_rooms": {
                "total_rooms": len(all_rooms),
                "total_capacity": total_room_capacity,
                "total_bookings": len(all_bookings),
                "total_hours_booked": round(total_room_booked_hours, 1),
                "average_occupancy_rate": min(avg_room_occupancy, 100.0),
                "peak_booking_time": fmt_hour_range(peak_room_hour)
            },
            "workspaces": {
                "total_desks": len(all_desks),
                "total_zones": len(zone_reports),
                "total_reservations": len(all_desk_res),
                "total_hours_reserved": round(total_desk_reserved_hours, 1),
                "average_utilization_rate": min(avg_desk_utilization, 100.0),
                "peak_reservation_time": fmt_hour_range(peak_desk_hour)
            }
        },
        "rooms": room_reports,
        "desks": desk_reports,
        "zones": zone_reports
    }


def generate_resource_utilization_excel(data: Dict[str, Any]) -> io.BytesIO:
    wb = Workbook()

    NAVY = "0F172A"
    ORANGE = "EA580C"
    LIGHT_BG = "F8FAFC"
    BORDER_COLOR = "CBD5E1"

    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Segoe UI", size=15, bold=True, color=NAVY)
    subtitle_font = Font(name="Segoe UI", size=10, italic=True, color="475569")
    kpi_num_font = Font(name="Segoe UI", size=14, bold=True, color=ORANGE)
    kpi_label_font = Font(name="Segoe UI", size=9, bold=True, color="64748B")
    cell_font = Font(name="Segoe UI", size=10)
    bold_cell_font = Font(name="Segoe UI", size=10, bold=True)

    header_fill = PatternFill(start_color=NAVY, end_color=NAVY, fill_type="solid")
    kpi_fill = PatternFill(start_color=LIGHT_BG, end_color=LIGHT_BG, fill_type="solid")
    zebra_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color=BORDER_COLOR),
        right=Side(style='thin', color=BORDER_COLOR),
        top=Side(style='thin', color=BORDER_COLOR),
        bottom=Side(style='thin', color=BORDER_COLOR)
    )

    meta = data["metadata"]
    summary = data["summary"]
    rm_summary = summary["meeting_rooms"]
    ws_summary = summary["workspaces"]

    # -------------------------------------------------------------
    # SHEET 1: Meeting Room Occupancy Rates
    # -------------------------------------------------------------
    ws1 = wb.active
    ws1.title = "Meeting Rooms Occupancy"
    ws1.views.sheetView[0].showGridLines = True

    ws1.cell(row=1, column=1, value="INTRASPHERE FACILITY MANAGEMENT").font = Font(name="Segoe UI", size=10, bold=True, color=ORANGE)
    ws1.cell(row=2, column=1, value="Meeting Room Occupancy & Capacity Utilization Report").font = title_font
    ws1.cell(
        row=3, column=1,
        value=f"Period: {meta['start_date']} to {meta['end_date']}  |  Floor: {meta['floor']}  |  Operating Basis: {meta['operating_hours_per_unit']} hrs/room  |  Generated: {meta['generated_at']}"
    ).font = subtitle_font

    # KPI Block
    kpis = [
        ("Total Rooms", rm_summary["total_rooms"]),
        ("Total Seating Capacity", rm_summary["total_capacity"]),
        ("Total Bookings", rm_summary["total_bookings"]),
        ("Total Booked Hours", f"{rm_summary['total_hours_booked']} hrs"),
        ("Avg Room Occupancy", f"{rm_summary['average_occupancy_rate']}%"),
        ("Peak Booking Window", rm_summary["peak_booking_time"]),
    ]
    for idx, (label, val) in enumerate(kpis, start=1):
        cell_l = ws1.cell(row=5, column=idx, value=label)
        cell_l.font = kpi_label_font
        cell_l.fill = kpi_fill
        cell_l.alignment = Alignment(horizontal="center", vertical="center")
        cell_l.border = thin_border

        cell_v = ws1.cell(row=6, column=idx, value=val)
        cell_v.font = kpi_num_font
        cell_v.fill = kpi_fill
        cell_v.alignment = Alignment(horizontal="center", vertical="center")
        cell_v.border = thin_border

    headers = [
        "Room Code", "Room Name", "Floor", "Location", "Capacity",
        "Total Bookings", "Hours Booked", "Operating Hours", "Occupancy Rate (%)", "Status"
    ]
    for col_idx, h in enumerate(headers, start=1):
        c = ws1.cell(row=8, column=col_idx, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center" if col_idx in (3, 5, 6, 7, 8, 9, 10) else "left", vertical="center")
        c.border = thin_border

    curr_row = 9
    for rm in data["rooms"]:
        row_vals = [
            rm["room_code"],
            rm["room_name"],
            f"Floor {rm['floor']}",
            rm["location"],
            f"{rm['capacity']} Seats",
            rm["total_bookings"],
            rm["total_hours_booked"],
            rm["operating_hours"],
            f"{rm['occupancy_rate_percent']}%",
            rm["status"]
        ]
        for col_idx, val in enumerate(row_vals, start=1):
            c = ws1.cell(row=curr_row, column=col_idx, value=val)
            c.font = bold_cell_font if col_idx in (1, 9) else cell_font
            c.alignment = Alignment(horizontal="center" if col_idx in (3, 5, 6, 7, 8, 9, 10) else "left", vertical="center")
            c.border = thin_border
            if (curr_row % 2) == 0:
                c.fill = zebra_fill

            if col_idx == 9:
                c.font = Font(name="Segoe UI", size=10, bold=True, color="2563EB" if rm["occupancy_rate_percent"] >= 50 else ORANGE)

        curr_row += 1

    for col in ws1.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws1.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # -------------------------------------------------------------
    # SHEET 2: Workspace & Desk Utilization
    # -------------------------------------------------------------
    ws2 = wb.create_sheet(title="Workspace Utilization")
    ws2.views.sheetView[0].showGridLines = True

    ws2.cell(row=1, column=1, value="INTRASPHERE FACILITY MANAGEMENT").font = Font(name="Segoe UI", size=10, bold=True, color=ORANGE)
    ws2.cell(row=2, column=1, value="Workspace & Desk Utilization Rate Report").font = title_font
    ws2.cell(
        row=3, column=1,
        value=f"Period: {meta['start_date']} to {meta['end_date']}  |  Floor: {meta['floor']}  |  Zone: {meta['zone']}  |  Operating Basis: {meta['operating_hours_per_unit']} hrs/desk"
    ).font = subtitle_font

    ws_kpis = [
        ("Total Workstations", ws_summary["total_desks"]),
        ("Total Active Zones", ws_summary["total_zones"]),
        ("Total Reservations", ws_summary["total_reservations"]),
        ("Total Hours Reserved", f"{ws_summary['total_hours_reserved']} hrs"),
        ("Avg Desk Utilization", f"{ws_summary['average_utilization_rate']}%"),
        ("Peak Usage Window", ws_summary["peak_reservation_time"]),
    ]
    for idx, (label, val) in enumerate(ws_kpis, start=1):
        cell_l = ws2.cell(row=5, column=idx, value=label)
        cell_l.font = kpi_label_font
        cell_l.fill = kpi_fill
        cell_l.alignment = Alignment(horizontal="center", vertical="center")
        cell_l.border = thin_border

        cell_v = ws2.cell(row=6, column=idx, value=val)
        cell_v.font = kpi_num_font
        cell_v.fill = kpi_fill
        cell_v.alignment = Alignment(horizontal="center", vertical="center")
        cell_v.border = thin_border

    # Section 1: Zone Summary Table
    ws2.cell(row=8, column=1, value="Zone Level Utilization Summary").font = Font(name="Segoe UI", size=12, bold=True, color=NAVY)
    zone_headers = ["Zone Name", "Floor", "Total Desks", "Total Reservations", "Hours Reserved", "Utilization Rate (%)"]
    for col_idx, h in enumerate(zone_headers, start=1):
        c = ws2.cell(row=9, column=col_idx, value=h)
        c.font = header_font
        c.fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
        c.alignment = Alignment(horizontal="center" if col_idx >= 2 else "left", vertical="center")
        c.border = thin_border

    z_row = 10
    for zn in data["zones"]:
        z_vals = [
            zn["zone_name"],
            f"Floor {zn['floor']}",
            zn["total_desks"],
            zn["total_reservations"],
            zn["total_hours_reserved"],
            f"{zn['utilization_rate_percent']}%"
        ]
        for col_idx, val in enumerate(z_vals, start=1):
            c = ws2.cell(row=z_row, column=col_idx, value=val)
            c.font = bold_cell_font if col_idx in (1, 6) else cell_font
            c.alignment = Alignment(horizontal="center" if col_idx >= 2 else "left", vertical="center")
            c.border = thin_border
            if (z_row % 2) == 0:
                c.fill = zebra_fill
        z_row += 1

    # Section 2: Desk-Level Breakdown Table
    d_start_row = z_row + 2
    ws2.cell(row=d_start_row, column=1, value="Individual Desk Utilization Details").font = Font(name="Segoe UI", size=12, bold=True, color=NAVY)

    desk_headers = [
        "Desk Code", "Desk Name", "Floor", "Zone", "Workspace Type",
        "Total Reservations", "Hours Reserved", "Operating Hours", "Utilization Rate (%)", "Status"
    ]
    for col_idx, h in enumerate(desk_headers, start=1):
        c = ws2.cell(row=d_start_row + 1, column=col_idx, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center" if col_idx in (3, 6, 7, 8, 9, 10) else "left", vertical="center")
        c.border = thin_border

    d_row = d_start_row + 2
    for dk in data["desks"]:
        row_vals = [
            dk["desk_code"],
            dk["desk_name"],
            f"Floor {dk['floor']}",
            dk["zone"],
            dk["workspace_type"],
            dk["total_reservations"],
            dk["total_hours_reserved"],
            dk["operating_hours"],
            f"{dk['utilization_rate_percent']}%",
            dk["status"]
        ]
        for col_idx, val in enumerate(row_vals, start=1):
            c = ws2.cell(row=d_row, column=col_idx, value=val)
            c.font = bold_cell_font if col_idx in (1, 9) else cell_font
            c.alignment = Alignment(horizontal="center" if col_idx in (3, 6, 7, 8, 9, 10) else "left", vertical="center")
            c.border = thin_border
            if (d_row % 2) == 0:
                c.fill = zebra_fill
        d_row += 1

    for col in ws2.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws2.column_dimensions[col_letter].width = max(max_len + 4, 12)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def generate_resource_utilization_pdf(data: Dict[str, Any]) -> io.BytesIO:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(letter),
        leftMargin=36,
        rightMargin=36,
        topMargin=36,
        bottomMargin=45
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=18,
        leading=22,
        textColor=colors.HexColor("#0F172A")
    )
    subtitle_style = ParagraphStyle(
        'DocSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=13,
        textColor=colors.HexColor("#475569")
    )
    section_title_style = ParagraphStyle(
        'SectionTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=11,
        leading=15,
        textColor=colors.HexColor("#0F172A")
    )
    th_style = ParagraphStyle(
        'TableHeader',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8,
        leading=10,
        textColor=colors.white,
        alignment=1
    )
    td_style = ParagraphStyle(
        'TableCell',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=7.5,
        leading=10,
        textColor=colors.HexColor("#1E293B")
    )
    td_center_style = ParagraphStyle(
        'TableCellCenter',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=7.5,
        leading=10,
        textColor=colors.HexColor("#1E293B"),
        alignment=1
    )
    kpi_num_style = ParagraphStyle(
        'KpiNum',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=13,
        leading=16,
        textColor=colors.HexColor("#EA580C"),
        alignment=1
    )
    kpi_label_style = ParagraphStyle(
        'KpiLabel',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=7.5,
        leading=9,
        textColor=colors.HexColor("#475569"),
        alignment=1
    )

    meta = data["metadata"]
    summary = data["summary"]
    rm_s = summary["meeting_rooms"]
    ws_s = summary["workspaces"]

    story = []
    story.append(Paragraph("<font color='#EA580C'><b>INTRASPHERE ENTERPRISE</b></font> • FACILITY MANAGEMENT", subtitle_style))
    story.append(Spacer(1, 3))
    story.append(Paragraph("Resource Utilization & Occupancy Rates Comprehensive Report", title_style))
    story.append(Spacer(1, 4))
    story.append(Paragraph(
        f"<b>Filter Period:</b> {meta['start_date']} to {meta['end_date']} ({meta['working_days']} working days) &nbsp;|&nbsp; "
        f"<b>Operating Basis:</b> {meta['operating_hours_per_unit']} hrs/unit &nbsp;|&nbsp; "
        f"<b>Floor:</b> {meta['floor']} &nbsp;|&nbsp; "
        f"<b>Zone:</b> {meta['zone']} &nbsp;|&nbsp; "
        f"<b>Generated On:</b> {meta['generated_at']}",
        subtitle_style
    ))
    story.append(Spacer(1, 14))

    # Executive KPI Summary Table
    kpi_data = [
        [
            Paragraph("FACILITY EFFICIENCY", kpi_label_style),
            Paragraph("ROOM OCCUPANCY", kpi_label_style),
            Paragraph("ROOM BOOKINGS", kpi_label_style),
            Paragraph("WORKSPACE UTILIZATION", kpi_label_style),
            Paragraph("DESK RESERVATIONS", kpi_label_style),
            Paragraph("PEAK ROOM HOURS", kpi_label_style),
        ],
        [
            Paragraph(f"{summary['overall_space_efficiency_score']}%", kpi_num_style),
            Paragraph(f"{rm_s['average_occupancy_rate']}%", kpi_num_style),
            Paragraph(str(rm_s['total_bookings']), kpi_num_style),
            Paragraph(f"{ws_s['average_utilization_rate']}%", kpi_num_style),
            Paragraph(str(ws_s['total_reservations']), kpi_num_style),
            Paragraph(rm_s['peak_booking_time'], kpi_num_style),
        ]
    ]
    kpi_table = Table(kpi_data, colWidths=[122] * 6)
    kpi_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor("#F8FAFC")),
        ('BOX', (0, 0), (-1, -1), 1, colors.HexColor("#E2E8F0")),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
    ]))
    story.append(kpi_table)
    story.append(Spacer(1, 16))

    # Section A: Meeting Room Occupancy Rates
    story.append(Paragraph("1. Meeting Room Occupancy Rates & Capacity Breakdown", section_title_style))
    story.append(Spacer(1, 6))

    rm_headers = [
        Paragraph("Room Code", th_style),
        Paragraph("Room Name", th_style),
        Paragraph("Floor", th_style),
        Paragraph("Location", th_style),
        Paragraph("Capacity", th_style),
        Paragraph("Bookings", th_style),
        Paragraph("Hours Booked", th_style),
        Paragraph("Occupancy Rate", th_style),
        Paragraph("Status", th_style),
    ]

    rm_col_widths = [75, 140, 60, 110, 65, 60, 75, 85, 65]
    rm_table_data = [rm_headers]

    for rm in data["rooms"]:
        rate_color = "#16A34A" if rm["occupancy_rate_percent"] >= 50 else ("#D97706" if rm["occupancy_rate_percent"] >= 20 else "#64748B")
        rm_table_data.append([
            Paragraph(rm["room_code"], td_style),
            Paragraph(rm["room_name"], td_style),
            Paragraph(f"Floor {rm['floor']}", td_center_style),
            Paragraph(rm["location"], td_style),
            Paragraph(f"{rm['capacity']} Seats", td_center_style),
            Paragraph(str(rm["total_bookings"]), td_center_style),
            Paragraph(f"{rm['total_hours_booked']}h", td_center_style),
            Paragraph(f"<font color='{rate_color}'><b>{rm['occupancy_rate_percent']}%</b></font>", td_center_style),
            Paragraph(rm["status"], td_center_style),
        ])

    rm_table = Table(rm_table_data, colWidths=rm_col_widths, repeatRows=1)
    rm_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#0F172A")),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
        ('TOPPADDING', (0, 0), (-1, -1), 4.5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4.5),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
    ]))
    story.append(rm_table)
    story.append(Spacer(1, 16))

    # Section B: Workspace & Zone Utilization
    story.append(Paragraph("2. Workspace Zone & Desk Utilization Breakdown", section_title_style))
    story.append(Spacer(1, 6))

    # Zone summary table
    zone_headers = [
        Paragraph("Zone Name", th_style),
        Paragraph("Floor", th_style),
        Paragraph("Total Desks", th_style),
        Paragraph("Total Reservations", th_style),
        Paragraph("Hours Reserved", th_style),
        Paragraph("Zone Utilization Rate", th_style),
    ]
    zone_col_widths = [190, 80, 95, 115, 115, 140]
    zone_table_data = [zone_headers]

    for zn in data["zones"]:
        z_rate_color = "#16A34A" if zn["utilization_rate_percent"] >= 50 else ("#D97706" if zn["utilization_rate_percent"] >= 20 else "#64748B")
        zone_table_data.append([
            Paragraph(zn["zone_name"], td_style),
            Paragraph(f"Floor {zn['floor']}", td_center_style),
            Paragraph(str(zn["total_desks"]), td_center_style),
            Paragraph(str(zn["total_reservations"]), td_center_style),
            Paragraph(f"{zn['total_hours_reserved']}h", td_center_style),
            Paragraph(f"<font color='{z_rate_color}'><b>{zn['utilization_rate_percent']}%</b></font>", td_center_style),
        ])

    zone_table = Table(zone_table_data, colWidths=zone_col_widths, repeatRows=1)
    zone_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#1E293B")),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
        ('TOPPADDING', (0, 0), (-1, -1), 4.5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4.5),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
    ]))
    story.append(zone_table)

    doc.build(story, canvasmaker=NumberedCanvas)
    buffer.seek(0)
    return buffer
